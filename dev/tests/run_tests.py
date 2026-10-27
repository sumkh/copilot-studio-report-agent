#!/usr/bin/env python3
"""End-to-end tests for the report-development skill scripts.

Usage:
  python run_tests.py [--keep DIR]

Needs Python 3.10+ with openpyxl (Pillow optional). Optional official-schema checks run when:
  RDL_XSD          = path to the RDL 2016/01 XSD (MS-RDL open specification, section 5.8)
  PBIR_SCHEMA_DIR  = folder mirroring github.com/microsoft/json-schemas (plus lxml/jsonschema installed)
"""
import argparse
import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.normpath(os.path.join(HERE, "..", "..", "skills", "report-development"))
SCRIPTS = os.path.join(SKILL, "scripts")
REFS = os.path.join(SKILL, "references")
PY = sys.executable
results = []


def run(script, *args, expect=0):
    p = subprocess.run([PY, os.path.join(SCRIPTS, script)] + [str(a) for a in args], capture_output=True, text=True)
    if expect is not None and p.returncode != expect:
        raise AssertionError("%s %s exited %d (expected %d)\nSTDOUT:\n%s\nSTDERR:\n%s" % (
            script, " ".join(map(str, args)), p.returncode, expect, p.stdout[-2000:], p.stderr[-2000:]))
    return p


def check(name, fn):
    try:
        fn()
        results.append((name, True, ""))
        print("PASS  " + name)
    except Exception as e:  # noqa: BLE001
        results.append((name, False, str(e)))
        print("FAIL  " + name + "\n      " + str(e).replace("\n", "\n      ")[:1500])


def load(p):
    with open(p, encoding="utf-8") as f:
        return json.load(f)


def save(p, obj):
    with open(p, "w", encoding="utf-8") as f:
        json.dump(obj, f, indent=2)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--keep")
    a = ap.parse_args()
    work = a.keep or tempfile.mkdtemp(prefix="rdskill_")
    os.makedirs(work, exist_ok=True)
    fx = os.path.join(work, "fixtures")
    subprocess.run([PY, os.path.join(HERE, "make_fixtures.py"), fx], check=True, capture_output=True)
    xlsx = os.path.join(fx, "Sample_Monthly_Sales_Report.xlsx")
    rdl_sample = os.path.join(fx, "Sample_Sales_By_Region.rdl")
    pbip_sample = os.path.join(fx, "Sample_Sales_Overview_pbip.zip")
    pbit = os.path.join(fx, "Sample_Sales_Model.pbit")
    csv_actual = os.path.join(fx, "actual_query_results.csv")
    xsd = os.environ.get("RDL_XSD")
    schema_dir = os.environ.get("PBIR_SCHEMA_DIR")
    sys.path.insert(0, HERE)
    sys.path.insert(1, SCRIPTS)

    def t_profile():
        out = os.path.join(work, "profile.json")
        run("xlsx_profile.py", xlsx, "--out", out)
        p = load(out)
        s = p["sheets"][0]
        assert s["table"]["header_row"] == 4, s["table"]
        assert len(s["columns"]) == 6
        assert s["row_classes"] == {"detail": 10, "subtotal": 3, "grand_total": 1}, s["row_classes"]
        assert s["group_column_candidates"] == ["Region"]
        assert s["styles"]["alt_row_background"] == "#F2F2F2"
        assert [c["format"] for c in s["columns"]][2:] == ["#,##0", "#,##0.00", "#,##0.00", "0.0%"]
        h = p["data_source_hints"]
        assert any(c.get("model_id") == "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee" for c in h["connections"])
        assert h["power_query"][0]["sources"][0]["args"][:2] == ["sql-prod-01.contoso.local", "SalesDW"]
        assert p["rdl_suggestions"]["style"]["header"]["background"] == "#1F4E78"
        assert [l["name"] for l in p["workbook"]["sensitivity_labels"]] == ["General"]
        assert "Region" in run("xlsx_profile.py", xlsx, "--summary").stdout
    check("xlsx_profile reads layout, styles, formulas and data-source clues", t_profile)

    def t_rdl_inspect():
        r = json.loads(run("rdl_inspect.py", rdl_sample).stdout)
        assert [d["kind"] for d in r["data_sources"]] == ["sql", "shared_data_source"]
        assert any("shared data source" in w for w in r["warnings"])
        ds = r["datasets"][0]
        assert len(ds["fields"]) == 6 and ds["query_parameters"][0]["name"] == "@StartDate"
        assert {p["name"] for p in r["parameters"]} == {"StartDate", "EndDate", "Region"}
        tx = [i for i in r["report_items"] if i["type"] == "Tablix"][0]
        assert len(tx["columns"]) == 6 and tx["columns"][0]["row_header"]
    check("rdl_inspect extracts sources, queries, parameters and layout", t_rdl_inspect)

    def t_pbi_inspect():
        r = json.loads(run("pbi_inspect.py", pbip_sample).stdout)
        assert r["exemplar"]["usable"] and r["exemplar"]["visual_types"] == ["card", "clusteredColumnChart", "slicer", "tableEx"]
        assert r["live_connection"]["model_id"] == "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"
        assert any(t["type"] == "custom" for t in r["themes"])
        r2 = json.loads(run("pbi_inspect.py", pbit).stdout)
        assert len(r2["semantic_model"]["tables"]) == 3 and not r2["exemplar"]["usable"]
        sql = [s for s in r2["data_sources"] if s.get("function") == "Sql.Database"]
        assert sql and sql[0]["args"][0].startswith("sql-prod-01.contoso.local"), sql
    check("pbi_inspect reads PBIR exemplars and template models", t_pbi_inspect)

    spec_p = os.path.join(work, "spec.json")

    def t_gates():
        shutil.copy(os.path.join(REFS, "spec-example-paginated.json"), spec_p)
        run("spec_tool.py", "validate", spec_p)
        run("rdl_build.py", spec_p, "--out", os.path.join(work, "x.rdl"), expect=3)
        run("spec_tool.py", "confirm", spec_p, "--gate", "data", "--reply", "yes but change the cost column", expect=4)
        run("spec_tool.py", "confirm", spec_p, "--gate", "spec", "--reply", "CONFIRM SPEC", expect=1)
        run("spec_tool.py", "confirm", spec_p, "--gate", "data", "--reply", "CONFIRM DATA")
        run("spec_tool.py", "confirm", spec_p, "--gate", "spec", "--reply", "CONFIRM SPEC")
        run("spec_tool.py", "check", spec_p)
        assert "Step 4 of 8" in run("spec_tool.py", "summary", spec_p, "--gate", "data").stdout
        assert "Step 5 of 8" in run("spec_tool.py", "summary", spec_p, "--gate", "spec").stdout
        unverified = os.path.join(work, "spec_unverified.json")
        s0 = load(os.path.join(REFS, "spec-example-paginated.json"))
        s0["data_sources"][0]["access_verified"] = False
        save(unverified, s0)
        p0 = run("spec_tool.py", "confirm", unverified, "--gate", "data", "--reply", "CONFIRM DATA", expect=1)
        assert "sign-in" in p0.stderr, p0.stderr
        drift = os.path.join(work, "spec_drift.json")
        s = load(spec_p)
        s["datasets"][0]["verified"]["rows"] = 11            # volatile: must not invalidate the gate
        save(drift, s)
        run("spec_tool.py", "check", drift)
        s["datasets"][0]["command_text"] += "\n-- edited"
        save(drift, s)
        p = run("spec_tool.py", "check", drift, expect=3)
        assert "datasets" in p.stdout
        d = run("spec_tool.py", "diff", spec_p, drift).stdout
        assert "gates to re-confirm: data" in d
    check("confirmation gates: refusal, ordering, change detection", t_gates)

    def t_rdl_table():
        out = os.path.join(work, "Monthly_Sales_Report_v1.rdl")
        run("rdl_build.py", spec_p, "--out", out)
        v = run("rdl_validate.py", out, "--json", *(["--xsd", xsd] if xsd else []))
        res = json.loads(v.stdout)
        assert not res["errors"], res
        if xsd:
            assert res["stats"].get("xsd") == "validated", res["stats"]
        # per-column colour expression and a page-break group with per-sheet names
        s = load(spec_p)
        s["layout"]["table"]["columns"][5]["color"] = '=IIf(Me.Value < 0, "Red", "Black")'
        s["layout"]["table"]["row_groups"][0].update({"page_break": True, "sheet_name": "=Fields!Region.Value"})
        s["layout"]["title"]["placement"] = "page_header"
        s["confirmations"] = {}
        p2 = os.path.join(work, "spec_variant.json")
        save(p2, s)
        run("spec_tool.py", "confirm", p2, "--gate", "data", "--reply", "CONFIRM DATA")
        run("spec_tool.py", "confirm", p2, "--gate", "spec", "--reply", "CONFIRM SPEC")
        out2 = os.path.join(work, "variant.rdl")
        run("rdl_build.py", p2, "--out", out2)
        res2 = json.loads(run("rdl_validate.py", out2, "--json", *(["--xsd", xsd] if xsd else [])).stdout)
        assert not res2["errors"], res2
        rdl = open(out2, encoding="utf-8").read()
        assert "Me.Value &lt; 0" in rdl and "<PageName>=Fields!Region.Value</PageName>" in rdl and "<PageHeader>" in rdl
    check("rdl_build table layout passes validation%s" % (" + XSD" if xsd else ""), t_rdl_table)

    def t_rdl_matrix():
        mp = os.path.join(work, "matrix.json")
        shutil.copy(os.path.join(HERE, "spec-matrix-test.json"), mp)
        run("spec_tool.py", "confirm", mp, "--gate", "data", "--reply", "CONFIRM DATA")
        run("spec_tool.py", "confirm", mp, "--gate", "spec", "--reply", "CONFIRM SPEC")
        out = os.path.join(work, "Sales_by_Month_Matrix_v1.rdl")
        run("rdl_build.py", mp, "--out", out)
        res = json.loads(run("rdl_validate.py", out, "--json", *(["--xsd", xsd] if xsd else [])).stdout)
        assert not res["errors"], res
        rdl = open(out, encoding="utf-8").read()
        assert '<QueryParameter Name="@Year">' in rdl and "<TablixCorner>" in rdl
    check("rdl_build matrix layout passes validation%s" % (" + XSD" if xsd else ""), t_rdl_matrix)

    def t_rdl_negative():
        src = open(os.path.join(work, "Monthly_Sales_Report_v1.rdl"), encoding="utf-8").read()
        bad = src.replace('<QueryParameter Name="StartDate">', '<QueryParameter Name="@StartDate">', 1)
        bad = bad.replace("=Fields!Cost.Value</Value>", "=Fields!Costs.Value</Value>", 1)
        bad = bad.replace("<TablixColumn><Width>62pt</Width></TablixColumn>", "", 1)
        bad = bad.replace("Data Source=pbiazure://api.powerbi.com/;", "Data Source=pbiazure://api.powerbi.com/;Password=x;", 1)
        p = os.path.join(work, "bad.rdl")
        open(p, "w", encoding="utf-8").write(bad)
        res = json.loads(run("rdl_validate.py", p, "--json", expect=1).stdout)
        text = " | ".join(res["errors"])
        for needle in ("must not start with '@'", "Costs", "5 columns but 6 leaf column members", "credentials"):
            assert needle in text, (needle, text)
        res2 = json.loads(run("rdl_validate.py", rdl_sample, "--json", expect=1).stdout)
        assert any("shared data source" in e for e in res2["errors"])
    check("rdl_validate catches broken definitions", t_rdl_negative)

    def t_pbip():
        pp = os.path.join(work, "pbi.json")
        shutil.copy(os.path.join(REFS, "spec-example-powerbi.json"), pp)
        run("spec_tool.py", "confirm", pp, "--gate", "data", "--reply", "CONFIRM DATA")
        run("spec_tool.py", "confirm", pp, "--gate", "spec", "--reply", "CONFIRM SPEC")
        out = os.path.join(work, "Regional_Sales_Dashboard_v1_pbip.zip")
        run("pbip_build.py", pp, "--exemplar", pbip_sample, "--out", out)
        model = os.path.join(work, "model.json")
        save(model, {"tables": {"Geography": {"columns": ["Region"], "measures": []}, "Product": {"columns": ["Category"], "measures": []},
                                "Sales": {"columns": [], "measures": ["Total Sales", "Total Units"]}}})
        res = json.loads(run("pbip_validate.py", out, "--model-schema", model, "--json").stdout)
        assert not res["errors"], res
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            report = json.loads(z.read("Regional_Sales_Dashboard.Report/definition/report.json"))
            assert "filterConfig" not in report                       # exemplar report filters dropped
            assert any(n.endswith("ContosoCorporate.json") for n in names)  # theme kept
            chart = [json.loads(z.read(n)) for n in names if n.endswith("visual.json")]
            col = [c for c in chart if c["visual"]["visualType"] == "clusteredColumnChart"][0]
            assert "objects" not in col["visual"] or "dataPoint" not in col["visual"]["objects"]  # selector formatting dropped
        bad_model = os.path.join(work, "model_bad.json")
        save(bad_model, {"tables": {"Geography": {"columns": ["Region"]}, "Product": {"columns": ["Category"]}, "Sales": {"measures": ["Total Sales"]}}})
        res2 = json.loads(run("pbip_validate.py", out, "--model-schema", bad_model, "--json", expect=1).stdout)
        assert any("Total Units" in e for e in res2["errors"])
        if schema_dir:
            import schema_check as sc
            bad = {k: v for k, v in sc.validate_pbir_files(sc.load_pbip_zip(out), schema_dir).items() if v}
            assert not bad, bad
    check("pbip_build clones the exemplar and passes validation%s" % (" + JSON schemas" if schema_dir else ""), t_pbip)

    def t_model():
        mp = os.path.join(work, "model.json")
        shutil.copy(os.path.join(REFS, "spec-example-semantic-model.json"), mp)
        run("spec_tool.py", "validate", mp)
        run("model_build.py", mp, "--exemplar", pbip_sample, "--out", os.path.join(work, "x.zip"), expect=3)
        assert "phase 1 of 2" in run("spec_tool.py", "summary", mp, "--gate", "data").stdout
        run("spec_tool.py", "confirm", mp, "--gate", "data", "--reply", "CONFIRM DATA")
        run("spec_tool.py", "confirm", mp, "--gate", "spec", "--reply", "CONFIRM SPEC")
        out = os.path.join(work, "Sales_Analytics_Model_v1_pbip.zip")
        schema_out = os.path.join(work, "model_schema.json")
        run("model_build.py", mp, "--exemplar", pbip_sample, "--out", out, "--schema-out", schema_out)
        res = json.loads(run("model_validate.py", out, "--json").stdout)
        assert not res["errors"], res
        assert res["stats"]["tables"] == 4 and res["stats"]["measures"] == 4 and res["stats"]["relationships"] == 3, res["stats"]
        assert sorted(load(schema_out)["tables"]) == ["Date", "Product", "Region", "Sales"]
        import tmdl
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            assert "Sales_Analytics_Model.SemanticModel/definition.pbism" in names
            assert "Sales_Analytics_Model.SemanticModel/DAXQueries/Model check.dax" in names
            files = {n: z.read(n).decode("utf-8") for n in names if n.endswith(".tmdl")}
            pbir = json.loads(z.read("Sales_Analytics_Model.Report/definition.pbir"))
            assert pbir["datasetReference"] == {"byPath": {"path": "../Sales_Analytics_Model.SemanticModel"}}
        parsed = tmdl.parse(files)
        sales = parsed["tables"]["Sales"]
        assert "Margin %" in sales["measures"] and "DIVIDE(" in sales["measures"]["Margin %"]["expression"]
        assert sales["partitions"][0]["source"].startswith("let")
        assert parsed["tables"]["Date"]["properties"]["dataCategory"] == "Time"
        if schema_dir:
            import schema_check as sc
            bad = {k: v for k, v in sc.validate_pbir_files(sc.load_pbip_zip(out), schema_dir).items() if v}
            assert not bad, bad
        # Tampered TMDL is caught
        tampered = os.path.join(work, "tampered.zip")
        with zipfile.ZipFile(out) as zin, zipfile.ZipFile(tampered, "w") as zout:
            for n in zin.namelist():
                data = zin.read(n)
                if n.endswith("relationships.tmdl"):
                    data = data.replace(b"toColumn: Region.'Region Key'", b"toColumn: Region.'Region Code'")
                if n.endswith("tables/Sales.tmdl"):
                    data = data.replace(b"SUM('Sales'[Cost])", b"SUM('Sales'[Costs])")
                zout.writestr(n, data)
        res2 = json.loads(run("model_validate.py", tampered, "--json", expect=1).stdout)
        text = " | ".join(res2["errors"])
        assert "Region Code" in text and "Costs" in text, text
        # Spec errors: mismatched relationship types and an unknown measure reference
        bad = load(os.path.join(REFS, "spec-example-semantic-model.json"))
        bad["relationships"][0]["to_column"] = "Region"
        bad["measures"].append({"table": "Sales", "name": "Broken", "expression": "[Total Salez] * 2"})
        bp = os.path.join(work, "model_bad.json")
        save(bp, bad)
        v = run("spec_tool.py", "validate", bp, expect=1).stdout
        assert "different data types" in v and "Total Salez" in v, v
        # Model diagram preview
        p = run("layout_preview.py", mp, "--out", os.path.join(work, "model.png")).stdout
        assert p.strip().endswith((".png", ".svg"))
    check("model_build makes a TMDL semantic model that passes validation%s" % (" + JSON schemas" if schema_dir else ""), t_model)

    def t_model_sql():
        mp = os.path.join(REFS, "spec-example-semantic-model.json")
        out = run("model_sql_check.py", mp, "--group", "Region[Region],Product[Category]", "--measures",
                  "Total Units,Total Sales,Margin %", "--date-from", "2026-09-01", "--date-to", "2026-09-30",
                  "--alias", "Category=Product Category").stdout
        for needle in ("SUM(f.[Units]) AS [Total Units]", "LEFT JOIN [dbo].[DimRegion] AS d1 ON f.[RegionKey] = d1.[RegionKey]",
                       "AS [Product Category]", "BETWEEN '2026-09-01' AND '2026-09-30'", "AS [Margin %]", "GROUP BY"):
            assert needle in out, (needle, out)
        s = load(mp)
        s["measures"].append({"table": "Sales", "name": "Sales LY", "expression": "CALCULATE([Total Sales], SAMEPERIODLASTYEAR('Date'[Date]))"})
        p2 = os.path.join(work, "model_ly.json")
        save(p2, s)
        out2 = run("model_sql_check.py", p2, "--group", "Date[Year]", "--measures", "Total Sales,Sales LY").stdout
        assert "YEAR(f.[OrderDate])" in out2 and "NOTE: [Sales LY] can't be recomputed" in out2, out2
    check("model_sql_check recomputes simple measures in SQL and flags the rest", t_model_sql)

    def t_reconcile():
        r = run("reconcile.py", "--sample", xlsx, "--actual", csv_actual, "--keys", "Region,Product Category",
                "--measures", "Units,Sales Amount,Cost", "--out", os.path.join(work, "recon.xlsx"), expect=2)
        res = json.loads(r.stdout)
        assert res["counts"]["mismatched_values"] == 1 and res["first_differences"][0]["key"] == ["EMEA", "Bikes"], res["counts"]
        assert res["totals"]["Sales Amount"]["difference"] == 1250.0
        import openpyxl
        wb = openpyxl.load_workbook(xlsx, data_only=True)
        ws = wb.active
        for row in range(5, ws.max_row + 1):
            ws.cell(row=row, column=4).number_format = "#,##0"
        ws.freeze_panes = None
        sim = os.path.join(work, "exported_sim.xlsx")
        wb.save(sim)
        res2 = json.loads(run("reconcile.py", "--sample", xlsx, "--actual", sim, "--keys", "Region,Product Category",
                              "--measures", "Units,Sales Amount,Cost").stdout)
        assert res2["status"] == "PASS"
        cmp_ = json.loads(run("xlsx_compare.py", xlsx, sim, "--json", expect=2).stdout)
        assert any(d["item"] == "number format of 'Sales Amount'" and d["severity"] == "major" for d in cmp_["differences"])
        assert any(d["severity"] == "expected" for d in cmp_["differences"])
    check("reconcile and xlsx_compare find number and format differences", t_reconcile)

    def t_preview():
        out = os.path.join(work, "preview.xlsx")
        run("layout_preview.py", spec_p, "--out", out, "--rows", csv_actual)
        cmp_ = json.loads(run("xlsx_compare.py", xlsx, out, "--json", expect=None).stdout)
        majors = [d for d in cmp_["differences"] if d["severity"] == "major"]
        assert not majors, majors
        run("layout_preview.py", os.path.join(work, "matrix.json"), "--out", os.path.join(work, "preview_matrix.xlsx"))
        p = run("layout_preview.py", os.path.join(work, "pbi.json"), "--out", os.path.join(work, "wire.png")).stdout
        assert p.strip().endswith((".png", ".svg"))
    check("layout_preview matches the sample's formatting", t_preview)

    import test_handover
    test_handover.add_checks(os.path.join(work, "handover"), check)

    failed = [r for r in results if not r[1]]
    print("\n%d passed, %d failed. Work folder: %s" % (len(results) - len(failed), len(failed), work))
    if not a.keep and not failed:
        shutil.rmtree(work, ignore_errors=True)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
