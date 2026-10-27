#!/usr/bin/env python3
"""Make a layout preview from a report spec, to show the user before they confirm the specification.

Usage:
  python layout_preview.py spec.json --out preview.xlsx [--rows sample_rows.csv]   (paginated spec)
  python layout_preview.py spec.json --out preview.png                             (Power BI spec; .svg if Pillow is missing)
  python layout_preview.py spec_model.json --out model.png                         (semantic model diagram)

Paginated: an Excel workbook laid out like the expected export (title rows, header,
detail rows, subtotals, grand total) using the spec's fonts, fills, borders, number formats
and column widths. Values are placeholders unless --rows gives a few real rows (CSV with
the dataset's field names or data_field names as headers). A second sheet says what the
preview is and isn't.
Power BI: a wireframe of each page with the visual boxes, types and titles.
"""
import argparse
import csv
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import reportkit as rk  # noqa: E402

BORDERS = {"Solid": "thin", "Dashed": "dashed", "Dotted": "dotted", "Double": "double"}


def excel_format(fmt):
    """Rough .NET -> Excel number format for the preview."""
    if not fmt:
        return "General"
    if fmt in ("d",):
        return "dd/mm/yyyy"
    if re.search(r"[dMy]", re.sub(r'"[^"]*"', "", fmt)) and not re.search(r"[0#]", fmt):
        return fmt.replace("M", "m").replace("tt", "AM/PM").replace("H", "h")
    return fmt


def placeholder(field_type, i, label):
    if field_type in ("Integer",):
        return 1200 + 37 * i
    if field_type in ("Float", "Decimal"):
        return 15234.5 + 812.25 * i
    if field_type == "DateTime":
        import datetime
        return datetime.date(2026, 9, 1 + i % 28)
    if field_type == "Boolean":
        return i % 2 == 0
    return "%s %s" % (label, chr(65 + i % 26))


def paginated_preview(spec, out, rows_csv=None):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    st_all = spec.get("style") or {}
    base = {"font_family": st_all.get("font_family") or "Calibri", "font_size_pt": st_all.get("font_size_pt") or 10}

    def style(part):
        s = dict(base)
        s.update({k: v for k, v in (st_all.get(part) or {}).items() if v is not None})
        return s

    def apply(c, s, fmt=None, align=None):
        c.font = Font(name=s.get("font_family"), size=s.get("font_size_pt"), bold=s.get("font_weight") == "Bold",
                      italic=s.get("font_style") == "Italic", underline="single" if s.get("underline") else None,
                      color=(rk.normalize_hex(s.get("color")) or "#000000")[1:])
        bg = rk.normalize_hex(s.get("background"))
        if bg:
            c.fill = PatternFill("solid", fgColor=bg[1:])
        bs = s.get("border_style")
        if bs and bs != "None":
            side = Side(style=BORDERS.get(bs, "thin") if (s.get("border_width_pt") or 1) < 2 or bs != "Solid" else "medium",
                        color=(rk.normalize_hex(s.get("border_color")) or "#000000")[1:])
            sides = s.get("border_sides") or ["top", "bottom", "left", "right"]
            c.border = Border(**{k: side for k in sides})
        h = (align or s.get("text_align") or "general").lower()
        c.alignment = Alignment(horizontal=h if h != "general" else None, vertical={"Middle": "center"}.get(s.get("vertical_align"), (s.get("vertical_align") or "bottom").lower()),
                                wrap_text=bool(s.get("wrap")))
        if fmt:
            c.number_format = excel_format(fmt)

    ds = {d["name"]: d for d in spec["datasets"]}[spec["layout"]["table"]["dataset"]]
    ftypes = {f["name"]: f.get("type", "String") for f in ds["fields"]}
    lay = spec["layout"]
    t = lay["table"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ((spec.get("excel_export") or {}).get("initial_sheet_name") or "Preview")[:31]
    real = []
    if rows_csv:
        with open(rows_csv, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            alias = {f_["data_field"]: f_["name"] for f_ in ds["fields"] if f_.get("data_field")}
            for i, r in enumerate(reader):
                if i >= 20:
                    break
                real.append({alias.get(k, k): v for k, v in r.items()})
    r = 1
    if t.get("mode", "table") == "table":
        cols = t["columns"]
        ncol = len(cols)
    else:
        cols = None
        months = ["Jan", "Feb", "Mar"]
        ncol = len(t["row_groups"]) + len(t["values"]) * (len(months) + (1 if t["column_group"].get("total") else 0))
    title = (lay.get("title") or {}).get("text")
    if title:
        ws.cell(row=r, column=1, value=title)
        apply(ws.cell(row=r, column=1), style("title"))
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        r += 1
    for line in lay.get("subtitle_lines") or []:
        text = line
        if line.startswith("="):
            text = re.sub(r'Format\(Parameters!(\w+)\.Value[^)]*\)', r'[\1]', line[1:])
            text = re.sub(r"Parameters!(\w+)\.Value", r"[\1]", text).replace('"', "").replace(" & ", "")
        ws.cell(row=r, column=1, value=text)
        apply(ws.cell(row=r, column=1), style("subtitle"))
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        r += 1
    if title or lay.get("subtitle_lines"):
        r += 1
    if cols is not None:
        for i, c in enumerate(cols, 1):
            ws.cell(row=r, column=i, value=c["header"])
            apply(ws.cell(row=r, column=i), style("header"), align=style("header").get("text_align") or "Center")
            ws.column_dimensions[get_column_letter(i)].width = rk.pt_to_excel_width(c["width_pt"])
        ws.row_dimensions[r].height = t.get("header_row_height_pt", 20)
        r += 1
        groups = t.get("row_groups") or []
        gfield = groups[0]["field"] if groups else None
        data = real or [{**{f: placeholder(ty, i, f) for f, ty in ftypes.items()}, **({gfield: "%s %s" % (gfield, "AB"[i // 2])} if gfield else {})}
                        for i in range(4)]
        alt = st_all.get("alt_row_background")

        def value(c, row):
            if c.get("field"):
                v = row.get(c["field"])
                n = rk_num(v)
                return n if n is not None and ftypes.get(c["field"]) in ("Integer", "Float", "Decimal") else v
            return 0.123 if "%" in (c.get("format") or "") else None

        def total_row(label, rows_, part, label_field):
            nonlocal r
            idx = next((i for i, c in enumerate(cols) if c.get("field") == label_field), 0)
            for i, c in enumerate(cols):
                v = None
                if c.get("aggregate") in ("Sum", "Count", "Avg", "Min", "Max") and c.get("field"):
                    nums = [rk_num(x.get(c["field"])) for x in rows_]
                    nums = [n for n in nums if n is not None]
                    v = {"Sum": sum(nums), "Count": len(nums), "Avg": sum(nums) / len(nums) if nums else None,
                         "Min": min(nums) if nums else None, "Max": max(nums) if nums else None}[c["aggregate"]]
                elif c.get("aggregate") == "Expression":
                    v = 0.123
                elif i == idx:
                    v = label
                cell = ws.cell(row=r, column=i + 1, value=v)
                apply(cell, style(part), fmt=c.get("format") if c.get("aggregate") else None, align=c.get("align") if c.get("aggregate") else None)
            r += 1
        blocks = []
        if gfield:
            for row in data:
                key = row.get(gfield)
                if not blocks or blocks[-1][0] != key:
                    blocks.append((key, []))
                blocks[-1][1].append(row)
        else:
            blocks = [(None, data)]
        n = 0
        for key, rows_ in blocks:
            for row in rows_:
                for i, c in enumerate(cols):
                    cell = ws.cell(row=r, column=i + 1, value=value(c, row))
                    s = style("detail")
                    if alt and n % 2 == 1 and not s.get("background"):
                        s = dict(s, background=alt)
                    apply(cell, s, fmt=c.get("format"), align=c.get("align"))
                n += 1
                r += 1
            if groups and groups[0].get("subtotal"):
                lbl = groups[0].get("subtotal_label") or "Total"
                lbl = "%s Total" % key if "Fields!" in lbl else lbl
                total_row(lbl, rows_, "subtotal", gfield)
        if (t.get("grand_total") or {}).get("show"):
            total_row((t.get("grand_total") or {}).get("label") or "Grand Total", data, "grand_total", None)
    else:
        rgs, vals, cg = t["row_groups"], t["values"], t["column_group"]
        months = ["Jan", "Feb", "Mar"] + ([cg.get("total_label") or "Total"] if cg.get("total") else [])
        for i, g in enumerate(rgs, 1):
            ws.cell(row=r, column=i, value=g.get("header") or g["field"])
            apply(ws.cell(row=r, column=i), style("header"))
            ws.column_dimensions[get_column_letter(i)].width = rk.pt_to_excel_width(g["width_pt"])
        col = len(rgs) + 1
        for mth in months:
            for v in vals:
                ws.cell(row=r, column=col, value=mth if len(vals) == 1 else "%s %s" % (mth, v["header"]))
                apply(ws.cell(row=r, column=col), style("header"), align="Center")
                ws.column_dimensions[get_column_letter(col)].width = rk.pt_to_excel_width(v["width_pt"])
                col += 1
        r += 1
        for gi in range(2):
            for di in range(2):
                for i, g in enumerate(rgs, 1):
                    apply(ws.cell(row=r, column=i, value="%s %s" % (g["field"], "AB"[gi] if i == 1 else "XY"[di])), style("detail"))
                col = len(rgs) + 1
                for _ in months:
                    for v in vals:
                        apply(ws.cell(row=r, column=col, value=1000 + 111 * gi + 37 * di), style("detail"), fmt=v.get("format"), align="Right")
                        col += 1
                r += 1
        if (t.get("grand_total") or {}).get("show"):
            apply(ws.cell(row=r, column=1, value=(t.get("grand_total") or {}).get("label") or "Grand Total"), style("grand_total"))
            for cc in range(2, col):
                apply(ws.cell(row=r, column=cc, value=4592 if cc > len(rgs) else None), style("grand_total"))
    pg = lay.get("page") or {}
    ws.page_setup.orientation = (pg.get("orientation") or "Landscape").lower()
    ws.sheet_view.showGridLines = False
    about = wb.create_sheet("About this preview")
    about["A1"] = "Layout preview for '%s' (spec version %s)" % (spec["report"]["name"], spec["report"].get("version", 1))
    about["A2"] = ("Numbers are placeholders, not real data." if not real else "Rows are a small sample from the test query.")
    about["A3"] = "The real report is produced by Power BI from the confirmed data source; totals are calculated there."
    about["A4"] = "Check: column order and headers, groups and totals, fonts, colours, number formats and widths."
    about.column_dimensions["A"].width = 110
    wb.save(out)


def rk_num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


def powerbi_preview(spec, out):
    fields = {f["id"]: "%s[%s]" % (f["table"], f["name"]) for f in spec.get("fields") or []}
    pages = spec.get("pages") or []
    scale = 0.75
    try:
        from PIL import Image, ImageDraw
        use_png = out.lower().endswith(".png")
    except ImportError:
        use_png = False
        if out.lower().endswith(".png"):
            out = out[:-4] + ".svg"
    if use_png:
        imgs = []
        for p in pages:
            w, h = int(p.get("width", 1280) * scale), int(p.get("height", 720) * scale)
            im = Image.new("RGB", (w, h + 30), "white")
            d = ImageDraw.Draw(im)
            d.text((8, 8), "Page: %s (wireframe)" % p["name"], fill="black")
            for v in p.get("visuals") or []:
                pos = v["position"]
                box = [int(pos["x"] * scale), int(pos["y"] * scale) + 30, int((pos["x"] + pos["width"]) * scale), int((pos["y"] + pos["height"]) * scale) + 30]
                d.rectangle(box, outline="#1F4E78", width=2, fill="#EEF3F9")
                d.text((box[0] + 6, box[1] + 6), "%s: %s" % (v["type"], v.get("title") or ""), fill="#1F4E78")
                y = box[1] + 22
                for role, ids in (v.get("roles") or {}).items():
                    d.text((box[0] + 6, y), "%s: %s" % (role, ", ".join(fields.get(i, i) for i in ids))[:80], fill="#333333")
                    y += 14
            imgs.append(im)
        total_h = sum(i.height for i in imgs) + 10 * (len(imgs) - 1)
        sheet = Image.new("RGB", (max(i.width for i in imgs), total_h), "#DDDDDD")
        y = 0
        for i in imgs:
            sheet.paste(i, (0, y))
            y += i.height + 10
        sheet.save(out)
        return out
    parts, y0 = [], 0
    for p in pages:
        w, h = p.get("width", 1280) * scale, p.get("height", 720) * scale
        parts.append('<rect x="0" y="%d" width="%d" height="%d" fill="white" stroke="#999"/>' % (y0, w, h + 30))
        parts.append('<text x="8" y="%d" font-family="Segoe UI" font-size="14">Page: %s (wireframe)</text>' % (y0 + 20, escape(p["name"])))
        for v in p.get("visuals") or []:
            pos = v["position"]
            x, y = pos["x"] * scale, pos["y"] * scale + 30 + y0
            parts.append('<rect x="%.0f" y="%.0f" width="%.0f" height="%.0f" fill="#EEF3F9" stroke="#1F4E78" stroke-width="2"/>' % (x, y, pos["width"] * scale, pos["height"] * scale))
            parts.append('<text x="%.0f" y="%.0f" font-family="Segoe UI" font-size="12" fill="#1F4E78">%s: %s</text>' % (x + 6, y + 16, escape(v["type"]), escape(v.get("title") or "")))
            for k, (role, ids) in enumerate((v.get("roles") or {}).items()):
                parts.append('<text x="%.0f" y="%.0f" font-family="Segoe UI" font-size="11" fill="#333">%s: %s</text>' % (x + 6, y + 32 + 14 * k, escape(role), escape(", ".join(fields.get(i, i) for i in ids))))
        y0 += h + 40
    width = max([p.get("width", 1280) * scale for p in pages] or [960])
    with open(out, "w", encoding="utf-8") as f:
        f.write('<svg xmlns="http://www.w3.org/2000/svg" width="%d" height="%d">%s</svg>' % (width, y0, "".join(parts)))
    return out


def escape(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def model_preview(spec, out):
    """Diagram of a semantic model spec: tables (fact in the middle), key columns, measures, relationships."""
    import tmdl
    tables = [dict(t) for t in spec.get("tables") or []]
    dt = spec.get("date_table") or {}
    if dt.get("create"):
        tables.append({"name": dt.get("name", "Date"), "role": "date",
                       "columns": [{"name": c[0], "hidden": c[3]} for c in tmdl.DATE_COLUMNS]})
    measures = {}
    for m in spec.get("measures") or []:
        measures.setdefault(m["table"], []).append(m["name"])
    facts = [t for t in tables if t.get("role") == "fact"] or tables[:1]
    dims = [t for t in tables if t not in facts]
    width, box_w, line_h = 1400, 300, 16

    def lines_for(t):
        cols = [c for c in t.get("columns") or [] if not c.get("hidden")][:8]
        hidden = len([c for c in t.get("columns") or [] if c.get("hidden")])
        out_ = ["%s (%s)" % (t["name"], t.get("role", "table"))] + ["  " + c["name"] for c in cols]
        if hidden:
            out_.append("  + %d hidden column%s" % (hidden, "s" if hidden > 1 else ""))
        for m in measures.get(t["name"], []):
            out_.append("  fx " + m)
        return out_
    pos, y = {}, 140
    for f in facts:
        pos[f["name"]] = (550, y, len(lines_for(f)) * line_h + 16)
        y += pos[f["name"]][2] + 50
    ly = ry = 60
    for i, d in enumerate(dims):
        h = len(lines_for(d)) * line_h + 16
        if i % 2 == 0:
            pos[d["name"]] = (60, ly, h)
            ly += h + 50
        else:
            pos[d["name"]] = (width - box_w - 60, ry, h)
            ry += h + 50
    height = max(y, ly, ry) + 40
    rels = [(r["from_table"], r["to_table"], r.get("active", True)) for r in spec.get("relationships") or []]
    rels += [(r["table"], dt.get("name", "Date"), r.get("active", True)) for r in dt.get("relationships") or []] if dt.get("create") else []

    def center(n):
        x, y_, h = pos[n]
        return x + box_w / 2, y_ + h / 2
    try:
        from PIL import Image, ImageDraw
        use_png = out.lower().endswith(".png")
    except ImportError:
        use_png = False
        if out.lower().endswith(".png"):
            out = out[:-4] + ".svg"
    title = "Model design: %s (spec v%s). Lines run from the many side (*) to the one side (1)." % (
        spec["report"]["name"], spec["report"].get("version", 1))
    if use_png:
        im = Image.new("RGB", (width, int(height)), "white")
        d = ImageDraw.Draw(im)
        d.text((20, 15), title, fill="black")
        for a, b, active in rels:
            if a in pos and b in pos:
                (x1, y1), (x2, y2) = center(a), center(b)
                d.line([(x1, y1), (x2, y2)], fill="#1F4E78" if active else "#999999", width=2)
        for t in tables:
            x, y_, h = pos[t["name"]]
            fill = "#DDEBF7" if t in facts else ("#FFF2CC" if t.get("role") == "date" else "#F2F2F2")
            d.rectangle([x, y_, x + box_w, y_ + h], outline="#1F4E78", width=2, fill=fill)
            for k, line in enumerate(lines_for(t)):
                d.text((x + 8, y_ + 8 + k * line_h), line[:46], fill="#1F4E78" if k == 0 else "#262626")
        for a, b, active in rels:          # cardinality markers on the visible part of each line
            if a in pos and b in pos:
                (x1, y1), (x2, y2) = center(a), center(b)
                d.text((x1 + (x2 - x1) * 0.38, y1 + (y2 - y1) * 0.38 - 14), "*", fill="black")
                d.text((x1 + (x2 - x1) * 0.62, y1 + (y2 - y1) * 0.62 - 14), "1" + ("" if active else " (inactive)"), fill="black")
        im.save(out)
        return out
    parts = ['<text x="20" y="25" font-family="Segoe UI" font-size="14">%s</text>' % escape(title)]
    for a, b, active in rels:
        if a in pos and b in pos:
            (x1, y1), (x2, y2) = center(a), center(b)
            parts.append('<line x1="%.0f" y1="%.0f" x2="%.0f" y2="%.0f" stroke="%s" stroke-width="2"/>' % (x1, y1, x2, y2, "#1F4E78" if active else "#999"))
    for t in tables:
        x, y_, h = pos[t["name"]]
        fill = "#DDEBF7" if t in facts else ("#FFF2CC" if t.get("role") == "date" else "#F2F2F2")
        parts.append('<rect x="%d" y="%d" width="%d" height="%d" fill="%s" stroke="#1F4E78" stroke-width="2"/>' % (x, y_, box_w, h, fill))
        for k, line in enumerate(lines_for(t)):
            parts.append('<text x="%d" y="%d" font-family="Segoe UI" font-size="12">%s</text>' % (x + 8, y_ + 20 + k * line_h, escape(line)))
    with open(out, "w", encoding="utf-8") as f:
        f.write('<svg xmlns="http://www.w3.org/2000/svg" width="%d" height="%d">%s</svg>' % (width, height, "".join(parts)))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("--out", required=True)
    ap.add_argument("--rows")
    a = ap.parse_args()
    spec = rk.load_json(a.spec)
    if spec["report_type"] == "paginated":
        paginated_preview(spec, a.out, a.rows)
        print("wrote", a.out)
    elif spec["report_type"] == "semantic_model":
        print("wrote", model_preview(spec, a.out))
    else:
        print("wrote", powerbi_preview(spec, a.out))


if __name__ == "__main__":
    main()
