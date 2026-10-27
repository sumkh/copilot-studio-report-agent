#!/usr/bin/env python3
"""Profile a sample Excel report: layout, formatting, formulas and data-source clues.

Usage:
  python xlsx_profile.py SAMPLE.xlsx [--out profile.json] [--summary] [--sheet NAME] [--no-samples]

The JSON output describes, per sheet, the main table (header row, columns, row
classes such as detail/subtotal/grand_total), its formatting (fonts, fills,
borders, number formats, column widths), page setup, and workbook-level clues
about where the data came from (Power Query sources, Analyze-in-Excel / OLAP
connections, pivot caches, external links). The `rdl_suggestions` block uses
report-spec vocabulary so it can be copied into a spec.

--summary prints a compact Markdown digest instead of JSON (no data values).
"""
import argparse
import base64
import collections
import datetime
import io
import os
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import reportkit as rk  # noqa: E402

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.styles.colors import COLOR_INDEX
except ImportError:  # pragma: no cover
    openpyxl = None

MAX_ROWS = 5000            # rows scanned per sheet
STYLE_SAMPLE_ROWS = 30     # detail rows sampled for styles
TOTAL_RE = re.compile(r"(?i)\b(grand\s*total|sub\s*-?\s*total|totals?)\b")
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
NS_S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
DEFAULT_THEME = ["#FFFFFF", "#000000", "#E7E6E6", "#44546A", "#4472C4", "#ED7D31",
                 "#A5A5A5", "#FFC000", "#5B9BD5", "#70AD47", "#0563C1", "#954F72"]
BORDER_MAP = {"thin": ("Solid", 1), "medium": ("Solid", 2), "thick": ("Solid", 3), "hair": ("Solid", 0.5),
              "dashed": ("Dashed", 1), "mediumDashed": ("Dashed", 2), "dotted": ("Dotted", 1),
              "dashDot": ("Dashed", 1), "mediumDashDot": ("Dashed", 2), "dashDotDot": ("Dotted", 1),
              "mediumDashDotDot": ("Dotted", 2), "slantDashDot": ("Dashed", 2), "double": ("Double", 3)}
ALIGN_MAP = {"left": "Left", "center": "Center", "centerContinuous": "Center", "right": "Right",
             "justify": "Left", "distributed": "Center", "fill": "General", "general": "General"}
VALIGN_MAP = {"top": "Top", "center": "Middle", "justify": "Middle", "distributed": "Middle", "bottom": "Bottom"}


# ----------------------------------------------------------------------------
# Workbook-level (zip) inspection
# ----------------------------------------------------------------------------

def theme_colors(z):
    names = sorted(n for n in z.namelist() if n.startswith("xl/theme/") and n.endswith(".xml"))
    if not names:
        return list(DEFAULT_THEME)
    try:
        root = ET.fromstring(z.read(names[0]))
        scheme = root.find(".//" + NS_A + "clrScheme")
        vals = {}
        for child in scheme:
            tag = child.tag.replace(NS_A, "")
            el = child[0]
            val = el.get("val") if el.tag.endswith("srgbClr") else el.get("lastClr")
            vals[tag] = "#" + val.upper() if val else None
        order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4",
                 "accent5", "accent6", "hlink", "folHlink"]
        return [vals.get(k) or DEFAULT_THEME[i] for i, k in enumerate(order)]
    except Exception:
        return list(DEFAULT_THEME)


def connections(z, warnings):
    out = []
    if "xl/connections.xml" not in z.namelist():
        return out
    try:
        root = ET.fromstring(z.read("xl/connections.xml"))
    except ET.ParseError as e:
        warnings.append("connections.xml unreadable: %s" % e)
        return out
    for c in root.iter(NS_S + "connection"):
        entry = {"id": c.get("id"), "name": c.get("name"), "description": c.get("description")}
        db = c.find(NS_S + "dbPr")
        if db is not None:
            conn, secret = rk.mask_secrets(db.get("connection") or "")
            props = rk.parse_conn_string(db.get("connection") or "")
            entry.update({"connection": conn, "command": db.get("command"), "command_type": db.get("commandType")})
            if secret:
                warnings.append("connection '%s' stores a password/secret - it will NOT be reused" % c.get("name"))
            if "mashup" in props.get("provider", "").lower():
                entry["kind"] = "power_query"
                entry["power_query_name"] = props.get("location")
            else:
                entry.update(rk.classify_connection(props.get("provider"), db.get("connection") or ""))
        web = c.find(NS_S + "webPr")
        if web is not None:
            entry.update({"kind": "web", "url": web.get("url")})
        txt = c.find(NS_S + "textPr")
        if txt is not None:
            entry.update({"kind": "text_file", "source_file": txt.get("sourceFile")})
        out.append(entry)
    return out


def pivot_caches(z):
    out = []
    for n in sorted(z.namelist()):
        if not re.match(r"xl/pivotCache/pivotCacheDefinition\d+\.xml$", n):
            continue
        try:
            root = ET.fromstring(z.read(n))
        except ET.ParseError:
            continue
        src = root.find(NS_S + "cacheSource")
        entry = {"part": n, "source_type": src.get("type") if src is not None else None,
                 "connection_id": src.get("connectionId") if src is not None else None}
        if src is not None:
            ws = src.find(NS_S + "worksheetSource")
            if ws is not None:
                entry["worksheet_source"] = {k: ws.get(k) for k in ("sheet", "ref", "name") if ws.get(k)}
        entry["fields"] = [cf.get("name") for cf in root.iter(NS_S + "cacheField")][:80]
        hier = []
        for h in root.iter(NS_S + "cacheHierarchy"):
            hier.append({"unique_name": h.get("uniqueName"), "caption": h.get("caption"),
                         "is_measure": h.get("measure") == "1"})
        if hier:
            entry["olap_hierarchies"] = hier[:120]
        out.append(entry)
    return out


def power_query(z, warnings):
    """Decode the DataMashup blob (MS-QDEFF) in customXml to list Power Query sources."""
    out = []
    for n in z.namelist():
        if not (n.startswith("customXml/item") and n.endswith(".xml")) or "Props" in n:
            continue
        text = rk.decode_text(z.read(n))
        if "DataMashup" not in text:
            continue
        m = re.search(r"<DataMashup[^>]*>(.*?)</DataMashup>", text, re.S)
        if not m:
            continue
        try:
            raw = base64.b64decode(re.sub(r"\s+", "", m.group(1)))
            pkg_len = int.from_bytes(raw[4:8], "little")
            with zipfile.ZipFile(io.BytesIO(raw[8:8 + pkg_len])) as pz:
                section = rk.decode_text(pz.read("Formulas/Section1.m"))
        except Exception as e:  # malformed or unknown version
            warnings.append("Power Query definitions found but could not be decoded: %s" % e)
            continue
        for qname, code in rk.m_shared_queries(section).items():
            out.append({"query": qname, "sources": rk.m_sources(code), "m_preview": code[:500]})
    return out


def external_links(z):
    out = []
    for n in z.namelist():
        if re.match(r"xl/externalLinks/_rels/externalLink\d+\.xml\.rels$", n):
            try:
                root = ET.fromstring(z.read(n))
            except ET.ParseError:
                continue
            for rel in root:
                if rel.get("Target"):
                    out.append(rel.get("Target"))
    return out


def sensitivity_labels(z):
    """Microsoft Purview (MSIP) labels stored as custom document properties."""
    labels = {}
    try:
        root = ET.fromstring(z.read("docProps/custom.xml"))
    except Exception:
        return []
    for prop in root:
        name = prop.get("name") or ""
        m = re.match(r"MSIP_Label_([0-9a-fA-F-]{36})_(Name|Enabled|Method)$", name)
        if m:
            labels.setdefault(m.group(1), {})[m.group(2).lower()] = (prop[0].text if len(prop) else None)
    return [{"id": k, "name": v.get("name"), "enabled": v.get("enabled"), "method": v.get("method")} for k, v in labels.items()]


def app_name(z):
    try:
        root = ET.fromstring(z.read("docProps/app.xml"))
        for el in root:
            if el.tag.endswith("Application"):
                return el.text
    except Exception:
        pass
    return None


# ----------------------------------------------------------------------------
# Cell / style helpers
# ----------------------------------------------------------------------------

def resolve_color(c, theme):
    if c is None:
        return None
    t = getattr(c, "type", None)
    try:
        if t == "rgb":
            return rk.normalize_hex(c.rgb if isinstance(c.rgb, str) else None)
        if t == "theme" and c.theme is not None and 0 <= c.theme < len(theme):
            return rk.apply_tint(theme[c.theme], c.tint or 0)
        if t == "indexed" and c.indexed is not None:
            if c.indexed == 64:
                return "#000000"
            if c.indexed == 65:
                return "#FFFFFF"
            if c.indexed < len(COLOR_INDEX):
                return rk.normalize_hex(COLOR_INDEX[c.indexed])
    except Exception:
        return None
    return None


def cell_style(cell, theme):
    f, fill, al, b = cell.font, cell.fill, cell.alignment, cell.border
    bg = None
    fill_type = getattr(fill, "fill_type", None)
    if fill_type and fill_type != "none":
        bg = resolve_color(getattr(fill, "fgColor", None), theme)   # solid and pattern fills
    sides = {}
    for side in ("left", "right", "top", "bottom"):
        s = getattr(b, side, None)
        if s is not None and s.style:
            sides[side] = (s.style, resolve_color(s.color, theme))
    return {
        "font_family": f.name if f else None,
        "font_size_pt": float(f.sz) if f and f.sz else None,
        "bold": bool(f.b) if f else False,
        "italic": bool(f.i) if f else False,
        "underline": bool(f.u) if f else False,
        "color": resolve_color(f.color, theme) if f else None,
        "background": bg,
        "borders": sides,
        "align": al.horizontal if al else None,
        "valign": al.vertical if al else None,
        "wrap": bool(al.wrap_text) if al else False,
        "indent": al.indent if al else 0,
        "number_format": cell.number_format,
    }


def dominant(styles):
    """Most common value per attribute across a list of style dicts."""
    if not styles:
        return {}
    keys = set().union(*styles)
    out = {}
    for k in keys:
        vals = [repr(s.get(k)) for s in styles]
        top = collections.Counter(vals).most_common(1)[0][0]
        out[k] = next(s.get(k) for s in styles if repr(s.get(k)) == top)
    return out


def to_rdl_style(st, default_font=None):
    if not st:
        return {}
    border = None
    sides = st.get("borders") or {}
    if sides:
        style, color = collections.Counter(sides.values()).most_common(1)[0][0]
        border_style, width = BORDER_MAP.get(style, ("Solid", 1))
        border = {"border_style": border_style, "border_color": color or "#000000", "border_width_pt": width,
                  "border_sides": sorted(sides)}
    out = {
        "font_family": st.get("font_family") or default_font,
        "font_size_pt": st.get("font_size_pt"),
        "font_weight": "Bold" if st.get("bold") else "Normal",
        "font_style": "Italic" if st.get("italic") else "Normal",
        "underline": bool(st.get("underline")),
        "color": st.get("color") or "#000000",
        "background": st.get("background"),
        "text_align": ALIGN_MAP.get(st.get("align") or "general", "General"),
        "vertical_align": VALIGN_MAP.get(st.get("valign") or "bottom", "Bottom"),
        "wrap": bool(st.get("wrap")),
    }
    if border:
        out.update(border)
    else:
        out["border_style"] = "None"
    return out


def kind_of(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, (int, float)):
        return "num"
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return "date"
    return "text"


def col_width_chars(ws, idx):
    letter = get_column_letter(idx)
    dims = ws.column_dimensions
    if letter in dims and dims[letter].width:
        return dims[letter].width, dims[letter].hidden
    for d in dims.values():
        if d.min and d.max and d.min <= idx <= d.max and d.width:
            return d.width, d.hidden
    default = ws.sheet_format.defaultColWidth if ws.sheet_format and ws.sheet_format.defaultColWidth else None
    return (default or 8.43), False


def translate_formula(formula, header_by_letter):
    def repl(m):
        letter = m.group(1)
        return "[%s]" % header_by_letter[letter] if letter in header_by_letter else m.group(0)
    f = re.sub(r"\$?([A-Z]{1,3})\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?", repl, formula)
    return f


def aggregate_from_formula(formula):
    f = (formula or "").upper()
    if re.search(r"SUBTOTAL\(\s*(9|109)\b", f) or "SUM(" in f:
        return "Sum"
    if re.search(r"SUBTOTAL\(\s*(1|101)\b", f) or "AVERAGE(" in f:
        return "Avg"
    if re.search(r"SUBTOTAL\(\s*(2|3|102|103)\b", f) or "COUNT" in f:
        return "Count"
    if "MIN(" in f:
        return "Min"
    if "MAX(" in f:
        return "Max"
    return None


# ----------------------------------------------------------------------------
# Table detection and sheet profiling
# ----------------------------------------------------------------------------

def read_grid(ws, max_rows):
    grid = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, max_rows), values_only=True):
        grid.append(list(row))
    return grid


def detect_table(grid, ws):
    best = None
    for r in range(min(60, len(grid))):
        row = grid[r]
        kinds = [kind_of(v) for v in row]
        cols = [i for i, k in enumerate(kinds) if k]
        texts = [i for i, k in enumerate(kinds) if k == "text"]
        if len(texts) < 2 or len(texts) < 0.6 * len(cols):
            continue
        c0, c1 = min(cols), max(cols)
        data_rows = numeric = 0
        for rr in range(r + 1, min(len(grid), r + 9)):
            seg = [kind_of(grid[rr][i]) if i < len(grid[rr]) else None for i in range(c0, c1 + 1)]
            if sum(1 for k in seg if k) >= max(1, 0.5 * len(texts)):
                data_rows += 1
            numeric += sum(1 for k in seg if k in ("num", "date"))
        if data_rows == 0:
            continue
        bold = sum(1 for i in texts if ws.cell(row=r + 1, column=i + 1).font.b)
        score = len(texts) * 2 + data_rows * 3 + min(numeric, 20) + bold
        if best is None or score > best[0]:
            best = (score, r, c0, c1)
    if not best:
        return None
    _, hr, c0, c1 = best
    last, empties = hr, 0
    for rr in range(hr + 1, len(grid)):
        row = grid[rr]
        if any(kind_of(row[i]) for i in range(c0, min(c1 + 1, len(row)))):
            last, empties = rr, 0
        else:
            empties += 1
            if empties >= 3:
                break
    return {"header_row": hr + 1, "first_col": c0 + 1, "last_col": c1 + 1,
            "data_first_row": hr + 2, "data_last_row": last + 1}


def classify_rows(grid, t, ws):
    c0, c1 = t["first_col"] - 1, t["last_col"] - 1
    rows = list(range(t["data_first_row"] - 1, t["data_last_row"]))
    text_cols = []
    for c in range(c0, c1 + 1):
        ks = [kind_of(grid[r][c]) for r in rows[:40] if c < len(grid[r])]
        ks = [k for k in ks if k]
        if ks and collections.Counter(ks).most_common(1)[0][0] == "text":
            text_cols.append(c)
    classes = {}
    for r in rows:
        vals = [grid[r][c] if c < len(grid[r]) else None for c in range(c0, c1 + 1)]
        if not any(kind_of(v) for v in vals):
            classes[r + 1] = "blank"
            continue
        texts = [str(grid[r][c]) for c in text_cols[:3] if c < len(grid[r]) and isinstance(grid[r][c], str)]
        joined = " ".join(texts)
        if TOTAL_RE.search(joined):
            classes[r + 1] = "grand_total" if re.search(r"(?i)grand", joined) else "subtotal"
            continue
        keys_empty = all(kind_of(grid[r][c] if c < len(grid[r]) else None) is None for c in text_cols)
        has_num = any(kind_of(v) == "num" for v in vals)
        is_bold = any(ws.cell(row=r + 1, column=c + 1).font.b for c in range(c0, c1 + 1))
        if text_cols and keys_empty and has_num and is_bold:
            classes[r + 1] = "subtotal"
        else:
            classes[r + 1] = "detail"
    non_blank = [r for r in sorted(classes) if classes[r] != "blank"]
    if non_blank and classes[non_blank[-1]] == "subtotal" and "grand_total" not in classes.values():
        label = " ".join(str(v) for v in grid[non_blank[-1] - 1] if isinstance(v, str)).strip().lower()
        n_sub = sum(1 for r in non_blank if classes[r] == "subtotal")
        if n_sub == 1 or label in ("total", "totals", "overall total", "report total", ""):
            classes[non_blank[-1]] = "grand_total"
    return classes, text_cols


def group_candidates(grid, t, classes, text_cols, headers):
    subs = [r for r, c in sorted(classes.items()) if c == "subtotal"]
    if not subs:
        return []
    segments, start = [], t["data_first_row"]
    for s in subs:
        segments.append([r for r in range(start, s) if classes.get(r) == "detail"])
        start = s + 1
    out = []
    for c in text_cols:
        ok = True
        for seg in segments:
            vals = {grid[r - 1][c] for r in seg if c < len(grid[r - 1]) and grid[r - 1][c] not in (None, "")}
            if len(vals) > 1:
                ok = False
                break
        if ok and any(segments):
            out.append(headers[c - (t["first_col"] - 1)])
    return out


def profile_sheet(ws, ws_f, theme, include_samples, warnings):
    sheet = {"name": ws.title, "state": ws.sheet_state, "max_row": ws.max_row, "max_column": ws.max_column}
    grid = read_grid(ws, MAX_ROWS)
    if ws.max_row and ws.max_row > MAX_ROWS:
        sheet["note"] = "profiled the first %d of %d rows" % (MAX_ROWS, ws.max_row)
    t = detect_table(grid, ws) if grid else None
    sheet["table"] = t
    feats = {
        "freeze_panes": ws.freeze_panes,
        "merged_ranges": [str(m) for m in ws.merged_cells.ranges][:100],
        "autofilter": ws.auto_filter.ref if ws.auto_filter else None,
        "charts": len(getattr(ws, "_charts", [])),
        "images": len(getattr(ws, "_images", [])),
        "pivot_tables": len(getattr(ws, "_pivots", [])),
        "data_validations": len(ws.data_validations.dataValidation) if ws.data_validations else 0,
        "show_gridlines": ws.sheet_view.showGridLines if ws.sheet_view else True,
        "tab_color": resolve_color(ws.sheet_properties.tabColor, theme) if ws.sheet_properties else None,
    }
    cf = []
    for rng in ws.conditional_formatting:
        for rule in rng.rules:
            cf.append({"range": str(rng.sqref), "type": rule.type, "operator": rule.operator,
                       "formula": list(rule.formula) if rule.formula else None})
    feats["conditional_formats"] = cf[:50]
    ps = ws.page_setup
    feats["page_setup"] = {
        "orientation": ps.orientation, "paper_size": ps.paperSize,
        "fit_to_width": ps.fitToWidth, "fit_to_height": ps.fitToHeight,
        "fit_to_page": bool(ws.sheet_properties.pageSetUpPr.fitToPage) if ws.sheet_properties.pageSetUpPr else False,
        "margins_in": {k: getattr(ws.page_margins, k) for k in ("left", "right", "top", "bottom")},
        "print_title_rows": ws.print_title_rows, "print_area": ws.print_area,
    }
    hf = {}
    for label, obj in (("header", ws.oddHeader), ("footer", ws.oddFooter)):
        for pos in ("left", "center", "right"):
            part = getattr(obj, pos, None)
            if part is not None and part.text:
                hf["%s_%s" % (label, pos)] = part.text
    feats["header_footer"] = hf
    sheet["features"] = feats
    if not t:
        return sheet

    c0, c1 = t["first_col"], t["last_col"]
    hr = t["header_row"]
    headers, seen = [], collections.Counter()
    for c in range(c0, c1 + 1):
        v = ws.cell(row=hr, column=c).value
        h = re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else "Column_%s" % get_column_letter(c)
        seen[h] += 1
        headers.append(h if seen[h] == 1 else "%s (%d)" % (h, seen[h]))
    header_by_letter = {get_column_letter(c): headers[c - c0] for c in range(c0, c1 + 1)}
    classes, text_cols = classify_rows(grid, t, ws)
    counts = collections.Counter(classes.values())
    sheet["row_classes"] = dict(counts)
    detail_rows = [r for r, c in sorted(classes.items()) if c == "detail"]
    sub_rows = [r for r, c in sorted(classes.items()) if c == "subtotal"]
    tot_rows = [r for r, c in sorted(classes.items()) if c == "grand_total"]
    sheet["subtotal_rows"] = sub_rows[:200]
    sheet["grand_total_rows"] = tot_rows

    # Title block above the header
    title = []
    merged = [m for m in ws.merged_cells.ranges]
    for r in range(1, hr):
        cells = []
        for c in range(1, max(c1, ws.max_column or c1) + 1):
            v = ws.cell(row=r, column=c).value
            if v not in (None, ""):
                rng = next((str(m) for m in merged if m.min_row <= r <= m.max_row and m.min_col <= c <= m.max_col), None)
                cells.append({"col": get_column_letter(c), "text": str(v)[:200], "merged": rng,
                              "style": to_rdl_style(cell_style(ws.cell(row=r, column=c), theme))})
        if cells:
            title.append({"row": r, "cells": cells})
    sheet["title_block"] = title
    multi_header = any(m.min_row == hr - 1 and m.max_col > m.min_col and m.min_col >= c0 for m in merged)
    if multi_header:
        sheet["multi_row_header"] = True

    # Columns
    columns = []
    sample_detail = detail_rows[:STYLE_SAMPLE_ROWS]
    for c in range(c0, c1 + 1):
        idx = c - c0
        vals = [grid[r - 1][c - 1] if c - 1 < len(grid[r - 1]) else None for r in detail_rows]
        kinds = [k for k in (kind_of(v) for v in vals) if k]
        top = collections.Counter(kinds).most_common(1)[0] if kinds else (None, 0)
        fmts = [ws.cell(row=r, column=c).number_format for r in sample_detail
                if kind_of(grid[r - 1][c - 1] if c - 1 < len(grid[r - 1]) else None)]
        fmt = collections.Counter(fmts).most_common(1)[0][0] if fmts else "General"
        dotnet, notes = rk.excel_numfmt_to_dotnet(fmt)
        dtype = top[0]
        if dtype == "num":
            nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
            has_decimals = bool(re.search(r"\.0", fmt or ""))
            dtype = "percent" if "%" in (fmt or "") else (
                "integer" if nums and not has_decimals and all(float(n).is_integer() for n in nums) else "decimal")
        if kinds and top[1] < 0.8 * len(kinds):
            dtype = "mixed"
        width, hidden = col_width_chars(ws, c)
        col = {
            "letter": get_column_letter(c), "header": headers[idx], "data_type": dtype or "empty",
            "non_empty": len(kinds), "distinct": len({repr(v) for v in vals if v not in (None, "")}),
            "excel_format": fmt, "format": dotnet, "format_notes": notes,
            "width_chars": round(width, 2), "width_pt": rk.excel_width_to_pt(width), "hidden": bool(hidden),
        }
        if include_samples:
            ex = []
            for v in vals:
                if v not in (None, "") and len(ex) < 3 and repr(v) not in [repr(e) for e in ex]:
                    ex.append(str(v)[:40] if not isinstance(v, (int, float)) else v)
            col["examples"] = ex
        # Formulas (formula workbook)
        if ws_f is not None:
            info = {}
            for cls, rows_ in (("detail", detail_rows), ("subtotal", sub_rows), ("grand_total", tot_rows)):
                fs = []
                for r in rows_[:500]:
                    fv = ws_f.cell(row=r, column=c).value
                    if isinstance(fv, str) and fv.startswith("="):
                        fs.append(fv)
                if fs:
                    info[cls] = {"count": len(fs), "example": fs[0], "readable": translate_formula(fs[0], header_by_letter)}
            if info:
                col["formulas"] = info
        agg = None
        f_tot = (col.get("formulas", {}).get("grand_total") or col.get("formulas", {}).get("subtotal") or {})
        if f_tot:
            agg = aggregate_from_formula(f_tot.get("example"))
        if agg is None and dtype in ("integer", "decimal") and tot_rows:
            detail_sum = sum(v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool))
            tv = grid[tot_rows[-1] - 1][c - 1] if c - 1 < len(grid[tot_rows[-1] - 1]) else None
            if isinstance(tv, (int, float)) and abs(detail_sum - tv) <= max(0.01, abs(tv) * 0.005):
                agg = "Sum"
        if dtype == "percent" and col.get("formulas", {}).get("detail"):
            agg = "Expression"
        col["aggregate_hint"] = agg
        columns.append(col)
    sheet["columns"] = columns

    # Styles by row class
    def styles_for(rows_):
        return [cell_style(ws.cell(row=r, column=c), theme) for r in rows_ for c in range(c0, c1 + 1)]
    header_st = dominant([cell_style(ws.cell(row=hr, column=c), theme) for c in range(c0, c1 + 1)])
    detail_st = dominant(styles_for(sample_detail))
    sub_st = dominant(styles_for(sub_rows[:10])) if sub_rows else {}
    tot_st = dominant(styles_for(tot_rows[:3])) if tot_rows else {}
    fills = [cell_style(ws.cell(row=r, column=c0), theme).get("background") for r in sample_detail]
    alt = None
    if len(fills) >= 4 and len(set(fills[:8])) == 2 and all(fills[i] == fills[i % 2] for i in range(len(fills[:8]))):
        alt = fills[1]
    per_col_align = {headers[c - c0]: ALIGN_MAP.get(dominant([cell_style(ws.cell(row=r, column=c), theme) for r in sample_detail]).get("align") or "general", "General")
                     for c in range(c0, c1 + 1)} if sample_detail else {}
    sheet["styles"] = {"header": header_st, "detail": detail_st, "subtotal": sub_st, "grand_total": tot_st,
                       "alt_row_background": alt}
    heights = {"header": ws.row_dimensions[hr].height,
               "detail": collections.Counter(ws.row_dimensions[r].height for r in sample_detail).most_common(1)[0][0] if sample_detail else None}
    sheet["row_heights_pt"] = heights
    sheet["group_column_candidates"] = group_candidates(grid, t, classes, text_cols, headers)

    # Fidelity notes for paginated Excel export
    notes = []
    if any(c.get("formulas") for c in columns):
        notes.append("The sample uses Excel formulas; paginated Excel exports contain values, not formulas.")
    if cf:
        notes.append("Conditional formatting rules become static colours; re-create them with Color/BackgroundColor expressions.")
    if feats["charts"]:
        notes.append("Charts export from paginated reports as pictures, not native Excel charts.")
    if feats["pivot_tables"]:
        notes.append("Pivot tables can't be reproduced; use a matrix (crosstab) layout instead.")
    if feats["autofilter"] or feats["data_validations"]:
        notes.append("AutoFilter and data validation aren't produced by paginated Excel exports.")
    data_merges = [m for m in merged if m.min_row > hr and m.min_col >= c0]
    if data_merges:
        notes.append("Merged cells inside the data area (%d) - the report will use unmerged cells." % len(data_merges))
    if feats["freeze_panes"]:
        notes.append("Frozen panes: paginated Excel exports lock only the page-header rows, not the table header.")
    if any(c["hidden"] for c in columns):
        notes.append("Hidden columns in the sample: %s - confirm whether to include them." % ", ".join(c["header"] for c in columns if c["hidden"]))
    if multi_header:
        notes.append("The sample has a two-row (grouped) header; a matrix layout may be needed.")
    sheet["fidelity_notes"] = notes
    sheet["_per_col_align"] = per_col_align
    return sheet


def rdl_suggestions(sheet, theme):
    t = sheet.get("table")
    if not t:
        return None
    st = sheet["styles"]
    font = st["detail"].get("font_family") or st["header"].get("font_family") or "Calibri"
    size = st["detail"].get("font_size_pt") or 10
    ps = sheet["features"]["page_setup"]
    landscape = (ps.get("orientation") == "landscape")
    name, w, h = rk.PAPER_SIZES_PT.get(ps.get("paper_size") or 9, rk.PAPER_SIZES_PT[9])
    if landscape:
        w, h = h, w
    margins = {k: round((v if v is not None else 0.5) * 72) for k, v in (ps.get("margins_in") or {}).items()}
    cols = []
    for c in sheet["columns"]:
        if c["data_type"] == "empty" and c["header"].startswith("Column_"):
            continue
        align = sheet["_per_col_align"].get(c["header"], "General")
        if align == "General":
            align = "Right" if c["data_type"] in ("integer", "decimal", "percent") else "Left"
        entry = {"header": c["header"], "width_pt": c["width_pt"], "format": c["format"], "align": align,
                 "data_type": c["data_type"], "aggregate_hint": c.get("aggregate_hint")}
        if c.get("format_notes"):
            entry["format_notes"] = c["format_notes"]
        f = c.get("formulas", {}).get("detail")
        if f:
            entry["formula_hint"] = f["readable"]
        cols.append(entry)
    title_lines = []
    title_style = None
    for row in sheet.get("title_block", []):
        title_lines.append(" ".join(cell["text"] for cell in row["cells"]))
        if title_style is None:
            title_style = row["cells"][0]["style"]
    return {
        "sheet": sheet["name"],
        "initial_sheet_name": sheet["name"][:31],
        "page": {"paper": name, "orientation": "Landscape" if landscape else "Portrait",
                 "width_pt": w, "height_pt": h, "margins_pt": margins},
        "title_lines": title_lines,
        "columns": cols,
        "style": {
            "font_family": font, "font_size_pt": size,
            "title": title_style or {},
            "header": to_rdl_style(st["header"], font),
            "detail": to_rdl_style(st["detail"], font),
            "alt_row_background": st.get("alt_row_background"),
            "subtotal": to_rdl_style(st["subtotal"], font) if st["subtotal"] else None,
            "grand_total": to_rdl_style(st["grand_total"], font) if st["grand_total"] else None,
        },
        "row_heights_pt": sheet.get("row_heights_pt"),
        "has_subtotals": bool(sheet.get("subtotal_rows")),
        "has_grand_total": bool(sheet.get("grand_total_rows")),
        "group_column_candidates": sheet.get("group_column_candidates"),
        "fidelity_notes": sheet.get("fidelity_notes"),
    }


def profile(path, only_sheet=None, include_samples=True):
    if openpyxl is None:
        rk.fail("openpyxl is not available in this environment; cannot profile .xlsx files")
    warnings = []
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        rk.fail("Legacy .xls files can't be read. Ask the user to save the workbook as .xlsx and upload again.")
    if not zipfile.is_zipfile(path):
        rk.fail("This workbook is encrypted or password-protected (for example by a sensitivity label with protection), "
                "so it can't be read. Ask the user for a copy they are allowed to share without protection, if policy permits.")
    with zipfile.ZipFile(path) as z:
        theme = theme_colors(z)
        conns = connections(z, warnings)
        pivots = pivot_caches(z)
        pq = power_query(z, warnings)
        links = external_links(z)
        application = app_name(z)
        labels = sensitivity_labels(z)
    wb = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False) if os.path.getsize(path) < 15 * 1024 * 1024 else None
    props = wb.properties
    result = {
        "file": os.path.basename(path),
        "workbook": {
            "sheets": wb.sheetnames, "application": application,
            "properties": {"title": props.title, "creator": props.creator, "last_modified_by": props.lastModifiedBy,
                           "created": props.created, "modified": props.modified},
            "defined_names": sorted(list(wb.defined_names.keys()))[:50] if hasattr(wb.defined_names, "keys") else [],
            "theme_colors": theme,
            "sensitivity_labels": labels,
        },
        "data_source_hints": {"connections": conns, "power_query": pq, "pivot_caches": pivots,
                              "external_links": links},
        "sheets": [],
        "warnings": warnings,
    }
    for name in wb.sheetnames:
        if only_sheet and name != only_sheet:
            continue
        ws = wb[name]
        if ws.sheet_state != "visible" and not only_sheet:
            result["sheets"].append({"name": name, "state": ws.sheet_state, "skipped": "hidden sheet"})
            continue
        if not hasattr(ws, "iter_rows") or ws.__class__.__name__ == "Chartsheet":
            continue
        ws_f = wb_f[name] if wb_f is not None else None
        result["sheets"].append(profile_sheet(ws, ws_f, theme, include_samples, warnings))
    primary = next((s for s in result["sheets"] if s.get("table")), None)
    result["primary_sheet"] = primary["name"] if primary else None
    result["rdl_suggestions"] = rdl_suggestions(primary, theme) if primary else None
    for s in result["sheets"]:
        s.pop("_per_col_align", None)
    return result


def summary_md(p):
    lines = ["# Excel sample profile: %s" % p["file"], ""]
    lines.append("Sheets: %s; primary table on **%s**." % (", ".join(p["workbook"]["sheets"]), p.get("primary_sheet")))
    labels = p["workbook"].get("sensitivity_labels") or []
    lines.append("Sensitivity label: %s" % (", ".join(l.get("name") or l["id"] for l in labels) if labels else "none found"))
    hints = p["data_source_hints"]
    src_lines = []
    for c in hints["connections"]:
        desc = c.get("kind") or "connection"
        detail = c.get("model_name") or c.get("model_id") or c.get("server") or c.get("power_query_name") or c.get("url") or ""
        src_lines.append("- Connection '%s': %s %s" % (c.get("name"), desc, detail))
    for q in hints["power_query"]:
        for s in q["sources"]:
            src_lines.append("- Power Query '%s': %s(%s)" % (q["query"], s["function"], ", ".join(map(str, s["args"][:3]))))
    for pc in hints["pivot_caches"]:
        meas = [h["unique_name"] for h in pc.get("olap_hierarchies", []) if h["is_measure"]][:8]
        src_lines.append("- Pivot cache (%s): %d fields%s" % (pc["source_type"], len(pc["fields"]),
                                                               ("; measures " + ", ".join(meas)) if meas else ""))
    for l in hints["external_links"]:
        src_lines.append("- External link: %s" % l)
    lines.append("")
    lines.append("## Data-source clues")
    lines.extend(src_lines or ["- None found in the workbook. Ask the user where the data comes from."])
    for s in p["sheets"]:
        if not s.get("table"):
            continue
        t = s["table"]
        lines += ["", "## Sheet '%s'" % s["name"],
                  "Table: header row %d, columns %s-%s, data rows %d-%d. Row classes: %s." % (
                      t["header_row"], s["columns"][0]["letter"], s["columns"][-1]["letter"],
                      t["data_first_row"], t["data_last_row"], s.get("row_classes"))]
        if s.get("title_block"):
            lines.append("Title rows: " + " | ".join(" ".join(c["text"] for c in r["cells"]) for r in s["title_block"]))
        lines += ["", "| Column | Type | Format (.NET) | Width pt | Aggregate | Formula |", "|---|---|---|---|---|---|"]
        for c in s["columns"]:
            f = c.get("formulas", {})
            ftxt = "; ".join("%s: %s" % (k, v["readable"]) for k, v in f.items())[:80]
            lines.append("| %s | %s | %s | %s | %s | %s |" % (c["header"], c["data_type"], c["format"] or "", c["width_pt"],
                                                         c.get("aggregate_hint") or "", ftxt))
        st = s["styles"]
        h = st["header"]
        lines.append("")
        lines.append("Header style: %s %spt%s, colour %s, fill %s." % (h.get("font_family"), h.get("font_size_pt"),
                     " bold" if h.get("bold") else "", h.get("color"), h.get("background")))
        d = st["detail"]
        lines.append("Detail style: %s %spt, colour %s, fill %s, alternate fill %s." % (d.get("font_family"), d.get("font_size_pt"),
                     d.get("color"), d.get("background"), st.get("alt_row_background")))
        if s.get("group_column_candidates"):
            lines.append("Subtotal grouping looks like it is by: %s." % ", ".join(s["group_column_candidates"]))
        ps = s["features"]["page_setup"]
        lines.append("Page: %s, paper %s, freeze %s." % (ps.get("orientation"), ps.get("paper_size"), s["features"].get("freeze_panes")))
        if s.get("fidelity_notes"):
            lines += ["", "Export fidelity notes:"] + ["- " + n for n in s["fidelity_notes"]]
    if p.get("warnings"):
        lines += ["", "Warnings:"] + ["- " + w for w in p["warnings"]]
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--out")
    ap.add_argument("--sheet")
    ap.add_argument("--summary", action="store_true")
    ap.add_argument("--no-samples", action="store_true")
    a = ap.parse_args()
    p = profile(a.xlsx, a.sheet, include_samples=not a.no_samples)
    if a.out:
        rk.dump_json(p, a.out)
    if a.summary:
        print(summary_md(p))
    elif not a.out:
        rk.dump_json(p)


if __name__ == "__main__":
    main()
