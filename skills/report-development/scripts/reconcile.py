#!/usr/bin/env python3
"""Reconcile the numbers in the user's sample Excel report with the new report's data.

Usage:
  python reconcile.py --sample SAMPLE.xlsx [--sample-sheet NAME]
                      --actual RESULTS.csv|EXPORTED.xlsx [--actual-sheet NAME]
                      --keys "Region,Product Category" --measures "Units,Sales Amount"
                      [--map "Sales Amount=[Sales Amount]" ...] [--tolerance-abs 0.01] [--tolerance-rel 0]
                      [--out reconciliation.xlsx] [--json]

--actual is either query output saved as CSV (pre-handoff check) or the Excel file the user
exported from the finished report (round-trip check). Subtotal and grand-total rows in Excel
files are detected and excluded from row matching; both sides are summed by the key columns
so different grains still compare. Printed totals in the sample are checked as well.
Prints a JSON summary (counts, totals, up to 20 differences). Detailed differences go to
--out, never to the chat. Exit code 0 = PASS, 2 = differences found.
"""
import argparse
import collections
import csv
import datetime
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import reportkit as rk  # noqa: E402

try:
    import openpyxl
    import xlsx_profile as xp
except ImportError:  # pragma: no cover
    openpyxl = None


def norm_header(h):
    h = str(h or "")
    h = re.sub(r"^.*\[(.+)\]$", r"\1", h.strip())     # Table[Column] or [Measure] -> Column
    return re.sub(r"[\s_]+", "", h).casefold()


def norm_key(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return re.sub(r"\s+", " ", str(v)).strip().casefold()


def to_num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        neg = s.startswith("(") and s.endswith(")")
        s = s.strip("()").rstrip("%")
        try:
            n = float(s)
            return -n if neg else n
        except ValueError:
            return None
    return None


def read_excel_table(path, sheet=None):
    if openpyxl is None:
        rk.fail("openpyxl is needed to read Excel files")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else next((w for w in wb.worksheets if xp.detect_table(xp.read_grid(w, xp.MAX_ROWS), w)), wb.worksheets[0])
    grid = xp.read_grid(ws, 200000)
    t = xp.detect_table(grid, ws)
    if not t:
        rk.fail("no table found on sheet '%s' of %s" % (ws.title, os.path.basename(path)))
    classes, _ = xp.classify_rows(grid, t, ws)
    c0, c1 = t["first_col"] - 1, t["last_col"] - 1
    headers = [str(grid[t["header_row"] - 1][c] or "").strip() for c in range(c0, c1 + 1)]
    detail, totals = [], {"subtotal": [], "grand_total": []}
    for r in range(t["data_first_row"], t["data_last_row"] + 1):
        row = {headers[i]: (grid[r - 1][c] if c < len(grid[r - 1]) else None) for i, c in enumerate(range(c0, c1 + 1))}
        cls = classes.get(r, "detail")
        if cls == "detail":
            detail.append(row)
        elif cls in totals:
            totals[cls].append(row)
    return {"sheet": ws.title, "headers": headers, "rows": detail, "totals": totals}


def read_csv_table(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    if not rows:
        rk.fail("empty CSV: %s" % path)
    headers = [h.strip() for h in rows[0]]
    return {"sheet": None, "headers": headers, "rows": [dict(zip(headers, r)) for r in rows[1:] if any(x.strip() for x in r)],
            "totals": {"subtotal": [], "grand_total": []}}


def map_columns(wanted, headers, explicit):
    mapping, missing = {}, []
    by_norm = {norm_header(h): h for h in headers}
    for w in wanted:
        if w in explicit:
            if explicit[w] not in headers:
                missing.append(w)
            else:
                mapping[w] = explicit[w]
        elif w in headers:
            mapping[w] = w
        elif norm_header(w) in by_norm:
            mapping[w] = by_norm[norm_header(w)]
        else:
            # 'Product Category' vs 'Product[Category]': accept a unique suffix/prefix match
            nw = norm_header(w)
            cands = [h for n, h in by_norm.items() if n and (nw.endswith(n) or n.endswith(nw) or nw.startswith(n))]
            if len(cands) == 1 and cands[0] not in mapping.values():
                mapping[w] = cands[0]
            else:
                missing.append(w)
    return mapping, missing


def aggregate(table, keys, measures, mapping):
    sums = collections.OrderedDict()
    counts = collections.Counter()
    non_numeric = collections.Counter()
    for row in table["rows"]:
        k = tuple(norm_key(row.get(mapping[c])) for c in keys)
        counts[k] += 1
        acc = sums.setdefault(k, {m: 0.0 for m in measures})
        for m in measures:
            v = row.get(mapping[m])
            n = to_num(v)
            if n is None:
                if v not in (None, ""):
                    non_numeric[m] += 1
            else:
                acc[m] += n
    labels = {}
    for row in table["rows"]:
        k = tuple(norm_key(row.get(mapping[c])) for c in keys)
        labels.setdefault(k, [str(row.get(mapping[c]) if row.get(mapping[c]) is not None else "") for c in keys])
    return sums, counts, labels, non_numeric


def within(a, b, tol_abs, tol_rel):
    return abs(a - b) <= max(tol_abs, tol_rel * abs(a))


def reconcile(a):
    sample = read_excel_table(a.sample, a.sample_sheet)
    actual = read_csv_table(a.actual) if a.actual.lower().endswith(".csv") else read_excel_table(a.actual, a.actual_sheet)
    keys = [k.strip() for k in a.keys.split(",") if k.strip()]
    measures = [m.strip() for m in a.measures.split(",") if m.strip()]
    explicit = dict(x.split("=", 1) for x in a.map or [])
    s_map, s_missing = map_columns(keys + measures, sample["headers"], {})
    a_map, a_missing = map_columns(keys + measures, actual["headers"], explicit)
    if s_missing or a_missing:
        rk.fail("column mapping failed. Not in sample: %s. Not in actual: %s. Actual headers: %s. Use --map 'Sample=Actual'."
                % (s_missing, a_missing, actual["headers"]))
    s_sum, s_cnt, s_lab, s_bad = aggregate(sample, keys, measures, s_map)
    a_sum, a_cnt, a_lab, a_bad = aggregate(actual, keys, measures, a_map)
    mismatches, missing, extra = [], [], []
    for k, sv in s_sum.items():
        if k not in a_sum:
            missing.append({"key": s_lab[k], **{m: sv[m] for m in measures}})
            continue
        for m in measures:
            if not within(sv[m], a_sum[k][m], a.tolerance_abs, a.tolerance_rel):
                mismatches.append({"key": s_lab[k], "measure": m, "sample": round(sv[m], 6), "actual": round(a_sum[k][m], 6),
                                   "difference": round(a_sum[k][m] - sv[m], 6)})
    for k, av in a_sum.items():
        if k not in s_sum:
            extra.append({"key": a_lab[k], **{m: av[m] for m in measures}})
    totals = {}
    for m in measures:
        s_total = sum(v[m] for v in s_sum.values())
        a_total = sum(v[m] for v in a_sum.values())
        printed = None
        if sample["totals"]["grand_total"]:
            printed = to_num(sample["totals"]["grand_total"][-1].get(s_map[m]))
        totals[m] = {"sample_detail_total": round(s_total, 6), "actual_total": round(a_total, 6),
                     "difference": round(a_total - s_total, 6), "match": within(s_total, a_total, a.tolerance_abs, a.tolerance_rel),
                     "sample_printed_grand_total": printed,
                     "sample_internally_consistent": None if printed is None else within(printed, s_total, a.tolerance_abs, a.tolerance_rel)}
    status = "PASS" if not (mismatches or missing or extra) and all(t["match"] for t in totals.values()) else "FAIL"
    notes = []
    if any(t["sample_internally_consistent"] is False for t in totals.values()):
        notes.append("The sample's printed grand total doesn't equal the sum of its own detail rows (manual edits or hidden rows?).")
    dup_s = sum(1 for c in s_cnt.values() if c > 1)
    dup_a = sum(1 for c in a_cnt.values() if c > 1)
    if dup_s or dup_a:
        notes.append("Rows were summed by key: %d sample keys and %d actual keys appear more than once." % (dup_s, dup_a))
    for side, bad in (("sample", s_bad), ("actual", a_bad)):
        for m, n in bad.items():
            notes.append("%d non-numeric %s values in '%s' were ignored." % (n, side, m))
    result = {"status": status, "sample": {"file": os.path.basename(a.sample), "sheet": sample["sheet"], "rows": len(sample["rows"]), "keys": len(s_sum)},
              "actual": {"file": os.path.basename(a.actual), "sheet": actual["sheet"], "rows": len(actual["rows"]), "keys": len(a_sum)},
              "column_mapping": {"sample": s_map, "actual": a_map}, "totals": totals,
              "counts": {"mismatched_values": len(mismatches), "missing_in_actual": len(missing), "extra_in_actual": len(extra)},
              "first_differences": mismatches[:20], "first_missing": missing[:10], "first_extra": extra[:10],
              "tolerance": {"absolute": a.tolerance_abs, "relative": a.tolerance_rel}, "notes": notes}
    if a.out:
        write_report(a.out, result, mismatches, missing, extra, keys, measures)
        result["report_file"] = a.out
    return result


def write_report(path, result, mismatches, missing, extra, keys, measures):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Status", result["status"]])
    ws.append(["Sample", "%s (%s), %d rows" % (result["sample"]["file"], result["sample"]["sheet"], result["sample"]["rows"])])
    ws.append(["Compared with", "%s, %d rows" % (result["actual"]["file"], result["actual"]["rows"])])
    ws.append(["Tolerance", "absolute %s, relative %s" % (result["tolerance"]["absolute"], result["tolerance"]["relative"])])
    ws.append([])
    ws.append(["Measure", "Sample total", "Actual total", "Difference", "Match", "Sample printed total"])
    for m, t in result["totals"].items():
        ws.append([m, t["sample_detail_total"], t["actual_total"], t["difference"], "yes" if t["match"] else "NO", t["sample_printed_grand_total"]])
    ws.append([])
    for n in result["notes"]:
        ws.append(["Note", n])
    s = wb.create_sheet("Differences")
    s.append(keys + ["Measure", "Sample", "Actual", "Difference"])
    for d in mismatches:
        s.append(d["key"] + [d["measure"], d["sample"], d["actual"], d["difference"]])
    for title, rows in (("Missing in new report", missing), ("Extra in new report", extra)):
        s = wb.create_sheet(title[:31])
        s.append(keys + measures)
        for d in rows:
            s.append(d["key"] + [d[m] for m in measures])
    mp = wb.create_sheet("Column mapping")
    mp.append(["Sample column", "Actual column"])
    for k, v in result["column_mapping"]["actual"].items():
        mp.append([k, v])
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sample", required=True)
    ap.add_argument("--sample-sheet")
    ap.add_argument("--actual", required=True)
    ap.add_argument("--actual-sheet")
    ap.add_argument("--keys", required=True)
    ap.add_argument("--measures", required=True)
    ap.add_argument("--map", action="append", help="'Sample header=Actual header', repeatable")
    ap.add_argument("--tolerance-abs", type=float, default=0.01)
    ap.add_argument("--tolerance-rel", type=float, default=0.0)
    ap.add_argument("--out")
    ap.add_argument("--json", action="store_true", help="(default) print the JSON summary")
    a = ap.parse_args()
    result = reconcile(a)
    print(json.dumps(result, indent=2, ensure_ascii=False, default=str))
    sys.exit(0 if result["status"] == "PASS" else 2)


if __name__ == "__main__":
    main()
